#!/usr/bin/env python3
"""Extract clinical table (S1) and the 21 triaptosis genes (S4A) from Braun 2020 Supplementary Tables."""
import csv
import openpyxl

P2 = "/Users/xiaolin/WorkBuddy/2026-09-04-20-00-08/paper2_triaptosis"
D = f"{P2}/data/external_validation/braun2020"
TRIA21 = {"ACTN2", "ANXA8", "ATG14", "ATP13A2", "BECN1", "CLCN4", "ELMO2", "FYCO1",
          "KEAP1", "KXD1", "MTM1", "NCKAP1", "NFE2L2", "NRBF2", "PIK3C3", "PIK3R4",
          "RAB9A", "SCARB2", "SH3GL3", "UVRAG", "WDR91"}

wb = openpyxl.load_workbook(f"{D}/41591_2020_839_MOESM2_ESM.xlsx", read_only=True)

rows = list(wb["S1_Clinical_and_Immune_Data"].iter_rows(min_row=2, values_only=True))
with open(f"{D}/braun2020_clinical.csv", "w", newline="") as f:
    w = csv.writer(f)
    for r in rows:
        if r[0] is not None:
            w.writerow(r)

ws = wb["S4A_RNA_Expression"]
it = ws.iter_rows(min_row=2, values_only=True)
header = next(it)
keep = [r for r in it if r[0] in TRIA21]
with open(f"{D}/braun2020_tria21_expr.csv", "w", newline="") as f:
    w = csv.writer(f)
    w.writerow(header)
    w.writerows(keep)
print("clinical rows", len(rows) - 1, "cols", len(rows[0]))
print("genes found", sorted(r[0] for r in keep), len(keep))
print("missing", sorted(TRIA21 - {r[0] for r in keep}))
