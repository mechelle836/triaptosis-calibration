#!/usr/bin/env python3
"""Write a Frontiers submission docx from manuscript_frontiers_oncology.md."""
import re
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.shared import Pt, Cm, RGBColor

ROOT = Path("/Users/xiaolin/WorkBuddy/2026-09-04-20-00-08/paper2_triaptosis")
SRC = ROOT / "manuscript_frontiers_oncology.md"
DST = ROOT / "submission_frontiers_oncology" / "Manuscript.docx"

INLINE = re.compile(
    r"(\*\*[^*]+\*\*|\*[^*]+\*|`[^`]+`|\^[^^]+\^)"
)


def add_runs(paragraph, text, size=12, base_italic=False, base_bold=False):
    for part in INLINE.split(text):
        if not part:
            continue
        italic, bold = base_italic, base_bold
        chunk = part
        if part.startswith("**") and part.endswith("**"):
            bold, chunk = True, part[2:-2]
        elif part.startswith("*") and part.endswith("*"):
            italic, chunk = True, part[1:-1]
        elif part.startswith("`") and part.endswith("`"):
            chunk = part[1:-1]
        elif part.startswith("^") and part.endswith("^"):
            run = paragraph.add_run(part[1:-1])
            run.font.size = Pt(size)
            run.font.name = "Times New Roman"
            run.italic = italic
            run.bold = bold
            run.font.superscript = True
            continue
        run = paragraph.add_run(chunk)
        run.font.size = Pt(size)
        run.font.name = "Times New Roman"
        run.italic = italic
        run.bold = bold


def add_para(doc, text, style=None, size=12, center=False, space_after=8, italic=False, bold=False):
    p = doc.add_paragraph()
    if style:
        p.style = style
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.line_spacing = 1.15
    if center:
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    add_runs(p, text, size=size, base_italic=italic, base_bold=bold)
    return p


def add_table(doc, rows):
    ncol = max(len(r) for r in rows)
    table = doc.add_table(rows=len(rows), cols=ncol)
    table.style = "Table Grid"
    for i, row in enumerate(rows):
        for j in range(ncol):
            cell = table.cell(i, j)
            cell.text = ""
            p = cell.paragraphs[0]
            val = row[j] if j < len(row) else ""
            add_runs(p, val, size=9, base_bold=(i == 0))
    doc.add_paragraph()


def main():
    lines = SRC.read_text(encoding="utf-8").splitlines()
    doc = Document()
    for section in doc.sections:
        section.top_margin = Cm(2.54)
        section.bottom_margin = Cm(2.54)
        section.left_margin = Cm(2.54)
        section.right_margin = Cm(2.54)
    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)
    style._element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")

    i = 0
    title_done = False
    while i < len(lines):
        line = lines[i].rstrip()
        if not line.strip() or line.strip() == "---":
            i += 1
            continue
        if line.startswith("|"):
            rows = []
            while i < len(lines) and lines[i].startswith("|"):
                raw = lines[i].strip().strip("|")
                cells = [c.strip() for c in raw.split("|")]
                if not all(re.fullmatch(r":?-{3,}:?", c.replace(" ", "")) or c == "" for c in cells):
                    rows.append(cells)
                i += 1
            if rows:
                add_table(doc, rows)
            continue
        if line.startswith("# "):
            add_para(doc, line[2:].strip(), size=16, center=True, space_after=12, bold=True)
            title_done = True
            i += 1
            continue
        if line.startswith("## "):
            add_para(doc, line[3:].strip(), size=14, space_after=8, bold=True)
            i += 1
            continue
        if line.startswith("### "):
            add_para(doc, line[4:].strip(), size=12, space_after=6, bold=True)
            i += 1
            continue
        # gather wrapped paragraph
        buf = [line.strip()]
        i += 1
        while i < len(lines) and lines[i].strip() and not lines[i].startswith(("#", "|", "---")):
            buf.append(lines[i].strip())
            i += 1
        text = " ".join(buf)
        center = (not title_done) or text.startswith("**Authors**") or text.startswith("**Affiliations**") or text.startswith("**Corresponding")
        add_para(doc, text, size=12, center=center and text.startswith("**Authors**") is False and not text.startswith("**Journal**"), space_after=8)
    DST.parent.mkdir(parents=True, exist_ok=True)
    doc.save(DST)
    print("wrote", DST, "paragraphs", len(doc.paragraphs), "tables", len(doc.tables))


if __name__ == "__main__":
    main()
